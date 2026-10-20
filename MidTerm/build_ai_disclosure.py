from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
OUTPUT = ROOT / "AI_Usage_Disclosure.docx"

document = Document()
section = document.sections[0]
section.top_margin = Inches(0.8)
section.bottom_margin = Inches(0.8)
section.left_margin = Inches(0.85)
section.right_margin = Inches(0.85)

normal = document.styles["Normal"]
normal.font.name = "Aptos"
normal.font.size = Pt(10.5)
normal.paragraph_format.space_after = Pt(8)
normal.paragraph_format.line_spacing = 1.08

title = document.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = title.add_run("AI Usage Disclosure")
run.bold = True
run.font.name = "Aptos Display"
run.font.size = Pt(22)
run.font.color.rgb = RGBColor(31, 78, 120)

subtitle = document.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = subtitle.add_run("FinTech 545 — Midterm Exercise")
run.bold = True
run.font.size = Pt(13)

document.add_paragraph()

details = document.add_table(rows=3, cols=2)
details.style = "Table Grid"
for row, (label, value) in enumerate(
    [
        ("Student", "____________________________"),
        ("Date", "October 19, 2026"),
        ("AI system", "OpenAI Codex"),
    ]
):
    details.cell(row, 0).text = label
    details.cell(row, 1).text = value
    details.cell(row, 0).paragraphs[0].runs[0].bold = True
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), "D9EAF7")
    details.cell(row, 0)._tc.get_or_add_tcPr().append(shading)

document.add_paragraph()

heading = document.add_paragraph()
run = heading.add_run("Summary of use")
run.bold = True
run.font.size = Pt(13)
run.font.color.rgb = RGBColor(31, 78, 120)

document.add_paragraph(
    "I used OpenAI Codex while completing this exercise. The AI assisted with "
    "writing and debugging Python code, checking numerical calculations, and "
    "formatting the computed results into a report."
)

document.add_paragraph("The assistance included:")
for item in [
    "Pairwise correlation, variance, eigenvalue, Higham-repair, and covariance calculations for Question 1.",
    "OLS market-model estimation and Normal VaR/ES calculations for Question 2.",
    "Marginal-distribution fitting, AICc comparison, copula fitting, simulation, and portfolio VaR/ES calculations for Question 3.",
    "Code execution checks and preparation of the final DOCX/PDF report.",
]:
    document.add_paragraph(item, style="List Bullet")

heading = document.add_paragraph()
run = heading.add_run("Representative requests")
run.bold = True
run.font.size = Pt(13)
run.font.color.rgb = RGBColor(31, 78, 120)

for prompt in [
    "Write Python code to calculate the pairwise correlation matrix and test whether it is positive semi-definite.",
    "Use Higham's method to obtain the nearest positive semi-definite correlation matrix.",
    "Estimate the market model and calculate 95% and 99% VaR and ES.",
    "Fit the requested marginal distributions and copulas, then simulate portfolio VaR and ES.",
    "Format the verified results as a PDF report.",
]:
    paragraph = document.add_paragraph(style="List Bullet")
    paragraph.add_run(prompt).italic = True

heading = document.add_paragraph()
run = heading.add_run("Responsibility statement")
run.bold = True
run.font.size = Pt(13)
run.font.color.rgb = RGBColor(31, 78, 120)

document.add_paragraph(
    "I reviewed the generated code and outputs and accept responsibility for "
    "the accuracy and interpretation of the submitted work. This document is "
    "a concise disclosure of AI use rather than a verbatim transcript of every interaction."
)

document.save(OUTPUT)
print(OUTPUT)
