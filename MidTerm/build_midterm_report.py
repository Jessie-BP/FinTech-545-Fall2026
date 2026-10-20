from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
DOCX_FILE = ROOT / "Midterm_Solutions.docx"

BLUE = "1F4E78"
LIGHT_BLUE = "D9EAF7"
LIGHT_GRAY = "F2F2F2"
WHITE = "FFFFFF"


def shade_cell(cell, fill):
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=60, start=70, bottom=60, end=70):
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        properties.append(margins)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def repeat_table_header(row):
    row_properties = row._tr.get_or_add_trPr()
    marker = OxmlElement("w:tblHeader")
    marker.set(qn("w:val"), "true")
    row_properties.append(marker)


def keep_row_together(row):
    row_properties = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    row_properties.append(cant_split)


def add_table(document, headers, rows, widths=None, font_size=8.5, first_column_left=True):
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.autofit = False
    header = table.rows[0]
    repeat_table_header(header)

    for column, value in enumerate(headers):
        cell = header.cells[column]
        cell.text = str(value)
        shade_cell(cell, BLUE)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_margins(cell)
        for paragraph in cell.paragraphs:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in paragraph.runs:
                run.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
                run.font.size = Pt(font_size)

    for row_number, values in enumerate(rows):
        row = table.add_row()
        keep_row_together(row)
        if row_number % 2 == 1:
            for cell in row.cells:
                shade_cell(cell, LIGHT_GRAY)
        for column, value in enumerate(values):
            cell = row.cells[column]
            cell.text = str(value)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            for paragraph in cell.paragraphs:
                paragraph.alignment = (
                    WD_ALIGN_PARAGRAPH.LEFT
                    if first_column_left and column == 0
                    else WD_ALIGN_PARAGRAPH.RIGHT
                )
                for run in paragraph.runs:
                    run.font.size = Pt(font_size)

    if widths:
        for row in table.rows:
            for column, width in enumerate(widths):
                row.cells[column].width = Inches(width)
    document.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_matrix(document, names, values, digits=6, font_size=7.3):
    headers = [""] + names
    rows = []
    for name, row in zip(names, values):
        rows.append([name] + [f"{value:.{digits}f}" for value in row])
    return add_table(
        document,
        headers,
        rows,
        widths=[0.55] + [(6.85 - 0.55) / len(names)] * len(names),
        font_size=font_size,
    )


def add_question(document, title, prompt):
    document.add_heading(title, level=1)
    paragraph = document.add_paragraph()
    paragraph.style = document.styles["Question Text"]
    run = paragraph.add_run(prompt)
    run.italic = True


def add_equation(document, text):
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(7)
    run = paragraph.add_run(text)
    run.font.name = "Cambria Math"
    run.font.size = Pt(11)
    return paragraph


document = Document()
section = document.sections[0]
section.top_margin = Inches(0.65)
section.bottom_margin = Inches(0.65)
section.left_margin = Inches(0.7)
section.right_margin = Inches(0.7)

styles = document.styles
normal = styles["Normal"]
normal.font.name = "Aptos"
normal.font.size = Pt(10)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.08

for name, size, color in (("Title", 25, BLUE), ("Heading 1", 17, BLUE), ("Heading 2", 12.5, BLUE)):
    style = styles[name]
    style.font.name = "Aptos Display"
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor.from_string(color)

styles["Heading 1"].paragraph_format.space_before = Pt(10)
styles["Heading 1"].paragraph_format.space_after = Pt(6)
styles["Heading 2"].paragraph_format.space_before = Pt(8)
styles["Heading 2"].paragraph_format.space_after = Pt(4)

question_style = styles.add_style("Question Text", WD_STYLE_TYPE.PARAGRAPH)
question_style.font.name = "Aptos"
question_style.font.size = Pt(9.5)
question_style.font.color.rgb = RGBColor(70, 70, 70)
question_style.paragraph_format.left_indent = Inches(0.18)
question_style.paragraph_format.right_indent = Inches(0.18)
question_style.paragraph_format.space_after = Pt(8)

for current_section in document.sections:
    header = current_section.header.paragraphs[0]
    header.text = "FinTech 545 — Quantitative Risk Management | Midterm Solutions"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in header.runs:
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor(100, 100, 100)

    footer = current_section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("Page ")
    run.font.size = Pt(8)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)

# Title page
document.add_paragraph().paragraph_format.space_after = Pt(55)
title = document.add_paragraph()
title.style = styles["Title"]
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.add_run("Midterm Exam Solutions")

subtitle = document.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = subtitle.add_run("FinTech 545 — Quantitative Risk Management")
run.bold = True
run.font.size = Pt(14)
run.font.color.rgb = RGBColor.from_string(BLUE)

document.add_paragraph().paragraph_format.space_after = Pt(26)
details = document.add_table(rows=3, cols=2)
details.alignment = 1
details.style = "Table Grid"
for left, right, row in (
    ("Student", "____________________________", 0),
    ("Date", "October 19, 2026", 1),
    ("Course", "FinTech 545", 2),
):
    details.cell(row, 0).text = left
    details.cell(row, 1).text = right
    shade_cell(details.cell(row, 0), LIGHT_BLUE)
    details.cell(row, 0).paragraphs[0].runs[0].bold = True
    for cell in details.rows[row].cells:
        set_cell_margins(cell, top=100, bottom=100)

document.add_paragraph().paragraph_format.space_after = Pt(25)
note = document.add_paragraph()
note.alignment = WD_ALIGN_PARAGRAPH.CENTER
note.add_run(
    "Arithmetic returns are used throughout. VaR and ES are reported as positive losses."
).italic = True

document.add_page_break()

# Question 1
add_question(
    document,
    "Question 1 — Pairwise Correlation with Missing Data (25 points)",
    "Using problem1.csv, calculate pairwise correlations; calculate each series' variance using all available observations; test whether the correlation matrix is positive semi-definite; if necessary, repair it with Higham's method; and combine the results into a covariance matrix.",
)

document.add_heading("1(a) Pairwise correlation matrix", level=2)
document.add_paragraph(
    "For each pair of series, I retained every row on which both members of the pair were observed and calculated the Pearson correlation on that pairwise-complete sample. Because the missing-value pattern differs across pairs, the effective sample is not identical for every entry."
)

q1_names = ["x1", "x2", "x3", "x4", "x5"]
q1_corr = [
    [1.000000, 0.332836, 0.418462, 0.255285, 0.732272],
    [0.332836, 1.000000, 0.251035, 0.287709, 0.787538],
    [0.418462, 0.251035, 1.000000, 0.483580, 0.684257],
    [0.255285, 0.287709, 0.483580, 1.000000, 0.782632],
    [0.732272, 0.787538, 0.684257, 0.782632, 1.000000],
]
add_matrix(document, q1_names, q1_corr)

document.add_heading("1(b) Variances using all available observations", level=2)
document.add_paragraph(
    "I calculated each sample variance independently, deleting missing values only from the series being measured. The denominator is nᵢ − 1. This preserves every valid observation in each individual series."
)
add_table(
    document,
    ["Series", "Valid observations", "Sample variance"],
    [
        ["x1", "41", "0.0004060739"],
        ["x2", "41", "0.0004800663"],
        ["x3", "36", "0.0001511102"],
        ["x4", "41", "0.0008417388"],
        ["x5", "42", "0.0002731496"],
    ],
    widths=[1.3, 2.0, 2.1],
)

document.add_heading("1(c) Positive semi-definite test", level=2)
document.add_paragraph(
    "A real symmetric matrix is positive semi-definite if and only if every eigenvalue is nonnegative. The eigenvalues of the pairwise correlation matrix are:"
)
add_equation(document, "[−0.08927867, 0.46133980, 0.71908517, 0.82343722, 3.08541647]")
document.add_paragraph(
    "The matrix is not positive semi-definite because its smallest eigenvalue is −0.08927867. This is a genuine negative value rather than numerical rounding. Pairwise deletion can produce this result because different correlations are estimated from different subsets of dates, so the entries need not form one internally consistent joint correlation structure."
)

document.add_heading("1(d) Higham nearest positive semi-definite correlation matrix", level=2)
document.add_paragraph(
    "I applied Higham's alternating-projection algorithm with Dykstra's correction. Each iteration projects onto the positive semi-definite cone by replacing negative eigenvalues with zero, then projects onto the unit-diagonal set by restoring every diagonal entry to one. The algorithm converged after 34 iterations."
)
q1_higham = [
    [1.000000, 0.349452, 0.425942, 0.270782, 0.694572],
    [0.349452, 1.000000, 0.260599, 0.307523, 0.739337],
    [0.425942, 0.260599, 1.000000, 0.492500, 0.662558],
    [0.270782, 0.307523, 0.492500, 1.000000, 0.737677],
    [0.694572, 0.739337, 0.662558, 0.737677, 1.000000],
]
add_matrix(document, q1_names, q1_higham)
document.add_paragraph("The corrected matrix has eigenvalues:")
add_equation(document, "[−1.62×10⁻¹², 0.45008153, 0.70347611, 0.80670847, 3.03973389]")
document.add_paragraph(
    "The first eigenvalue is zero to numerical precision; the displayed −1.62×10⁻¹² is floating-point error. Therefore, the corrected matrix is positive semi-definite."
)
add_table(
    document,
    ["Entry", "Original", "Higham", "Absolute change"],
    [
        ["(x2, x5)", "0.787538", "0.739337", "0.048201"],
        ["(x4, x5)", "0.782632", "0.737677", "0.044955"],
        ["(x1, x5)", "0.732272", "0.694572", "0.037700"],
        ["(x3, x5)", "0.684257", "0.662558", "0.021699"],
        ["(x2, x4)", "0.287709", "0.307523", "0.019814"],
    ],
    widths=[1.4, 1.4, 1.4, 1.6],
)
document.add_paragraph(
    "The entries involving x5 changed the most. The largest change was the correlation between x2 and x5, followed by the correlations between x4 and x5 and between x1 and x5."
)

document.add_heading("1(e) Covariance matrix", level=2)
document.add_paragraph(
    "Let D be the diagonal matrix of full-sample standard deviations from part (b), and let Rᴴ be the Higham correlation matrix from part (d). I formed the covariance matrix as Σ = D Rᴴ D. Consequently, each off-diagonal covariance is Rᴴᵢⱼσᵢσⱼ, while the diagonal exactly reproduces the variances from part (b)."
)
q1_cov = [
    [0.0004060739, 0.0001542911, 0.0001055115, 0.0001583111, 0.0002313238],
    [0.0001542911, 0.0004800663, 0.0000701891, 0.0001954867, 0.0002677277],
    [0.0001055115, 0.0000701891, 0.0001511102, 0.0001756475, 0.0001346081],
    [0.0001583111, 0.0001954867, 0.0001756475, 0.0008417388, 0.0003537164],
    [0.0002313238, 0.0002677277, 0.0001346081, 0.0003537164, 0.0002731496],
]
add_matrix(document, q1_names, q1_cov, digits=10, font_size=6.8)

document.add_page_break()

# Question 2
add_question(
    document,
    "Question 2 — VaR and ES for a Single Stock (20 points)",
    "Using problem2.csv, estimate rXYZ = α + βrSPY + ε for a holding of 100 XYZ shares. Assume arithmetic returns of SPY and ε are normally distributed. Calculate 95% and 99% VaR and ES in percent and explain the method.",
)

document.add_heading("2(a) VaR at 95% and 99%", level=2)
document.add_paragraph(
    "I first converted the 501 prices to 500 arithmetic daily returns. I estimated the market model by ordinary least squares. The fitted coefficients and distribution inputs are:"
)
add_table(
    document,
    ["Quantity", "Estimate"],
    [
        ["α", "−0.00030402"],
        ["β", "1.19331076"],
        ["SPY mean", "0.00025477"],
        ["SPY sample standard deviation", "0.01097207"],
        ["Residual sample standard deviation", "0.01459005"],
    ],
    widths=[3.3, 2.0],
)
document.add_paragraph(
    "Because rSPY and ε are assumed to be independent normal variables, their linear combination is also normal. The implied XYZ mean and standard deviation are:"
)
add_equation(document, "μXYZ = α + βμSPY = 0.00000000")
add_equation(document, "σXYZ = √(β²σ²SPY + σ²ε) = 0.01960353")
document.add_paragraph(
    "For confidence level c, I used the lower return quantile and reported the loss as a positive number: VaRᶜ = −[μXYZ + σXYZ Φ⁻¹(1−c)]. The number of shares affects dollar VaR but cancels when the result is expressed as a percentage of position value."
)
add_table(
    document,
    ["Confidence level", "Return cutoff", "VaR"],
    [
        ["95%", "−3.2245%", "3.2245%"],
        ["99%", "−4.5605%", "4.5605%"],
    ],
    widths=[1.8, 1.8, 1.5],
)

document.add_heading("2(b) ES at 95% and 99%", level=2)
document.add_paragraph(
    "Expected shortfall is the average loss conditional on the return falling below the corresponding VaR cutoff. For a normal return, with p = 1 − c and zₚ = Φ⁻¹(p), I used:"
)
add_equation(document, "ESᶜ = −μXYZ + σXYZ φ(zₚ)/p")
add_table(
    document,
    ["Confidence level", "VaR", "Expected shortfall"],
    [
        ["95%", "3.2245%", "4.0436%"],
        ["99%", "4.5605%", "5.2248%"],
    ],
    widths=[1.8, 1.6, 2.0],
)
document.add_paragraph(
    "ES is larger than VaR at both confidence levels because it averages losses beyond the quantile rather than reporting only the boundary of the tail. The increase from 95% to 99% also reflects the more extreme loss region being measured."
)

document.add_page_break()

# Question 3
add_question(
    document,
    "Question 3 — Copula VaR and ES for a Two-Book Portfolio (75 points)",
    "Using problem3.csv and problem3_holdings.csv, fit Normal, Student's t, and NIG marginal distributions to ten stock returns; compare Gaussian and Student's t copulas; and use the selected copula with at least 5,000 simulations to calculate 95% absolute dollar VaR and ES for each stock, each book, and the total portfolio.",
)

document.add_heading("3(a) Marginal distribution fitting and selection", level=2)
document.add_paragraph(
    "I converted the 1,001 prices into 1,000 arithmetic daily returns for each stock. Each Normal, Student's t, and normal inverse Gaussian (NIG) distribution was fitted by maximum likelihood. I compared models with the small-sample corrected Akaike information criterion:"
)
add_equation(document, "AICc = −2ℓ + 2k + 2k(k+1)/(n−k−1)")
document.add_paragraph(
    "Here k equals 2 for the Normal, 3 for Student's t, and 4 for NIG. Lower AICc is preferred because it rewards likelihood while charging more flexible distributions for their extra parameters."
)

normal_rows = [
    ["S1", "0.00090370", "0.01182025", "−6033.994"],
    ["S2", "0.00080026", "0.01533991", "−5512.706"],
    ["S3", "0.00042361", "0.01888852", "−5096.513"],
    ["S4", "0.00077999", "0.00992765", "−6382.974"],
    ["S5", "0.00070877", "0.01894165", "−5090.895"],
    ["S6", "0.00080469", "0.01507758", "−5547.203"],
    ["S7", "0.00069273", "0.01366672", "−5743.694"],
    ["S8", "0.00037400", "0.01710306", "−5295.107"],
    ["S9", "0.00174602", "0.02063642", "−4919.506"],
    ["S10", "0.00016480", "0.01305470", "−5835.325"],
]
document.add_paragraph("Normal distribution fits:").runs[0].bold = True
add_table(document, ["Stock", "μ", "σ", "AICc"], normal_rows, widths=[1.0, 1.6, 1.6, 1.5])

t_rows = [
    ["S1", "355368542.589", "0.00090366", "0.01182027", "−6031.982"],
    ["S2", "4.857877", "0.00058139", "0.01167261", "−5627.360"],
    ["S3", "4.160482", "0.00229110", "0.01365721", "−5238.463"],
    ["S4", "339036291.501", "0.00078000", "0.00992769", "−6380.962"],
    ["S5", "6.289248", "0.00083326", "0.01557655", "−5150.869"],
    ["S6", "3.823756", "0.00244489", "0.01065489", "−5688.740"],
    ["S7", "319103.363", "0.00069242", "0.01366808", "−5741.682"],
    ["S8", "3.483289", "0.00065411", "0.01142390", "−5493.288"],
    ["S9", "4.385214", "0.00338341", "0.01531014", "−5036.762"],
    ["S10", "4.121270", "0.00040472", "0.00942699", "−5974.857"],
]
document.add_paragraph("Student's t distribution fits:").runs[0].bold = True
add_table(document, ["Stock", "ν", "Location", "Scale", "AICc"], t_rows, widths=[0.8, 1.4, 1.4, 1.4, 1.3], font_size=8.0)

nig_rows = [
    ["S1", "2086.485600", "−588.204192", "0.25747326", "0.07655696", "−6030.215"],
    ["S2", "73.206598", "4.863823", "0.01659306", "−0.00030457", "−5622.063"],
    ["S3", "79.836689", "−36.344624", "0.01965684", "0.01047387", "−5281.641"],
    ["S4", "3368.189374", "−984.957946", "0.29028696", "0.08954907", "−6379.120"],
    ["S5", "73.576752", "−2.232562", "0.02595840", "0.00149685", "−5146.210"],
    ["S6", "91.682513", "−41.686207", "0.01462753", "0.00827200", "−5737.395"],
    ["S7", "1635.963518", "−443.712607", "0.27261055", "0.07748191", "−5739.888"],
    ["S8", "50.622840", "−2.886825", "0.01416134", "0.00118304", "−5487.094"],
    ["S9", "63.159483", "−22.926741", "0.02151486", "0.01012760", "−5065.510"],
    ["S10", "76.685731", "−5.368403", "0.01259531", "0.00104879", "−5970.702"],
]
document.add_paragraph("NIG distribution fits:").runs[0].bold = True
add_table(document, ["Stock", "α", "β", "δ", "μ", "AICc"], nig_rows, widths=[0.65, 1.15, 1.15, 1.15, 1.15, 1.15], font_size=7.5)

selection_rows = [
    ["S1", "−6033.994", "−6031.982", "−6030.215", "Normal"],
    ["S2", "−5512.706", "−5627.360", "−5622.063", "Student's t"],
    ["S3", "−5096.513", "−5238.463", "−5281.641", "NIG"],
    ["S4", "−6382.974", "−6380.962", "−6379.120", "Normal"],
    ["S5", "−5090.895", "−5150.869", "−5146.210", "Student's t"],
    ["S6", "−5547.203", "−5688.740", "−5737.395", "NIG"],
    ["S7", "−5743.694", "−5741.682", "−5739.888", "Normal"],
    ["S8", "−5295.107", "−5493.288", "−5487.094", "Student's t"],
    ["S9", "−4919.506", "−5036.762", "−5065.510", "NIG"],
    ["S10", "−5835.325", "−5974.857", "−5970.702", "Student's t"],
]
document.add_paragraph("AICc comparison and selected margins:").runs[0].bold = True
add_table(
    document,
    ["Stock", "Normal", "Student's t", "NIG", "Selected"],
    selection_rows,
    widths=[0.8, 1.25, 1.35, 1.25, 1.4],
    font_size=8.0,
)
document.add_paragraph(
    "Normal is selected for S1, S4, and S7. Their fitted t degrees of freedom are extremely large, so Student's t has effectively collapsed to a Normal and its extra parameter is not justified. Student's t is selected for S2, S5, S8, and S10, where finite degrees of freedom capture symmetric heavy tails. NIG is selected for S3, S6, and S9, where its additional skewness and tail flexibility produce the largest AICc improvement."
)

document.add_heading("3(b) Gaussian and Student's t copulas", level=2)
document.add_paragraph(
    "For each selected margin, I transformed observed returns to uniforms with its fitted CDF. I then fitted each copula by maximum likelihood. For the Gaussian copula, the correlation estimate is the sample correlation matrix of the normal scores Φ⁻¹(u). For the Student's t copula, I estimated the correlation matrix from the t scores t_ν⁻¹(u) at each trial ν and maximized the copula log likelihood over ν in [2.01, 100]; this profile approach approximates the joint maximum likelihood estimate of the correlation matrix and ν. Each copula estimates the same 45 correlations, and the Student's t copula adds ν, so I use k = 45 for the Gaussian and k = 46 for the Student's t in AICc and BIC. The Gaussian copula correlation matrix is:"
)

stock_names = [f"S{i}" for i in range(1, 11)]
copula_corr = [
    [1.0000, 0.5484, 0.5344, 0.5606, 0.5454, 0.2938, 0.2504, 0.2444, 0.2712, 0.2995],
    [0.5484, 1.0000, 0.5664, 0.5278, 0.5188, 0.3253, 0.2472, 0.2529, 0.2932, 0.2998],
    [0.5344, 0.5664, 1.0000, 0.5315, 0.5742, 0.3139, 0.3093, 0.2688, 0.3275, 0.3083],
    [0.5606, 0.5278, 0.5315, 1.0000, 0.5596, 0.3035, 0.2440, 0.2688, 0.2816, 0.2924],
    [0.5454, 0.5188, 0.5742, 0.5596, 1.0000, 0.3330, 0.2803, 0.2977, 0.2764, 0.2993],
    [0.2938, 0.3253, 0.3139, 0.3035, 0.3330, 1.0000, 0.5455, 0.5279, 0.5490, 0.5628],
    [0.2504, 0.2472, 0.3093, 0.2440, 0.2803, 0.5455, 1.0000, 0.5299, 0.5406, 0.5433],
    [0.2444, 0.2529, 0.2688, 0.2688, 0.2977, 0.5279, 0.5299, 1.0000, 0.5227, 0.5229],
    [0.2712, 0.2932, 0.3275, 0.2816, 0.2764, 0.5490, 0.5406, 0.5227, 1.0000, 0.5423],
    [0.2995, 0.2998, 0.3083, 0.2924, 0.2993, 0.5628, 0.5433, 0.5229, 0.5423, 1.0000],
]
add_matrix(document, stock_names, copula_corr, digits=4, font_size=6.1)
document.add_paragraph(
    "The minimum eigenvalue is 0.37844635, so the matrix is positive definite and requires no repair. The Student's t copula correlation matrix at the fitted ν is nearly identical (every entry within 0.001 of the Gaussian matrix, minimum eigenvalue 0.37806972), so the comparison below is driven by the likelihood and the extra parameter ν rather than by different dependence estimates."
)
add_table(
    document,
    ["Copula", "Log likelihood", "AICc", "BIC", "ν"],
    [
        ["Gaussian", "2113.924", "−4133.508", "−3916.999", "—"],
        ["Student's t", "2111.601", "−4126.665", "−3905.445", "99.999997"],
    ],
    widths=[1.4, 1.5, 1.3, 1.3, 1.2],
)
document.add_paragraph(
    "The Gaussian copula fits better because it has both the lower AICc and the lower BIC. The fitted t degrees of freedom reached the upper search boundary of approximately 100, which is close to the Gaussian limit. Therefore, the data do not provide enough evidence that an additional common heavy-tail dependence parameter improves the copula. I use the Gaussian copula in part (c)."
)

document.add_heading("3(c) Copula simulation, VaR, and ES", level=2)
document.add_paragraph(
    "I generated 50,000 draws from the fitted Gaussian copula using random seed 545. I transformed each simulated uniform through the inverse CDF of its selected marginal distribution. Each position was valued at its last observed price, and its one-day profit and loss was calculated as position value multiplied by the simulated arithmetic return. For each stock, book, and the total portfolio, 95% absolute VaR is the negative fifth percentile of P&L, and ES is the negative mean P&L in observations at or below that percentile."
)
risk_rows = [
    ["S1", "$48,329.48", "$898.54", "$1,132.97"],
    ["S2", "$46,697.43", "$1,069.61", "$1,556.43"],
    ["S3", "$24,267.28", "$803.61", "$1,200.53"],
    ["S4", "$43,596.81", "$678.47", "$855.20"],
    ["S5", "$37,727.45", "$1,094.17", "$1,521.92"],
    ["S6", "$46,326.64", "$1,212.25", "$1,846.59"],
    ["S7", "$40,144.31", "$878.66", "$1,101.91"],
    ["S8", "$26,045.88", "$648.89", "$1,031.66"],
    ["S9", "$98,248.71", "$3,365.61", "$5,161.64"],
    ["S10", "$24,512.38", "$477.48", "$706.66"],
    ["Book A", "$200,618.44", "$3,615.35", "$4,886.64"],
    ["Book B", "$235,277.92", "$5,461.00", "$7,933.73"],
    ["Total portfolio", "$435,896.36", "$7,640.46", "$10,648.60"],
]
risk_table = add_table(
    document,
    ["Position", "Current value", "95% VaR", "95% ES"],
    risk_rows,
    widths=[1.7, 1.7, 1.5, 1.5],
    font_size=8.5,
)
for row_index in (11, 12, 13):
    for cell in risk_table.rows[row_index].cells:
        shade_cell(cell, LIGHT_BLUE)
        for run in cell.paragraphs[0].runs:
            run.bold = True

document.add_paragraph(
    "S9 is the largest stand-alone risk source, with 95% VaR of $3,365.61 and ES of $5,161.64. This reflects both its large $98,248.71 position value and its NIG return margin, which allows skewness and heavier downside behavior. Book B consequently has more risk than Book A: Book B's VaR is $5,461.00 versus $3,615.35 for Book A, and its ES is $7,933.73 versus $4,886.64."
)
document.add_paragraph(
    "Risk is not additive. The total portfolio VaR of $7,640.46 is 31.34% below the sum of the ten stand-alone stock VaRs. Total ES of $10,648.60 is 33.92% below the sum of the stand-alone ES values. Correlations are generally around 0.52–0.57 within each book but only about 0.24–0.33 across books. Therefore, combining the two books provides meaningful diversification even though all positions are long and all pairwise dependence is positive."
)
document.add_paragraph(
    "ES exceeds VaR at the stock, book, and portfolio levels because ES averages the more severe losses beyond the VaR threshold. The gap is especially visible for the heavy-tailed or skewed margins, such as S9, and is why ES is useful alongside VaR when evaluating tail severity."
)

document.add_page_break()
document.add_heading("Reproducibility", level=1)
document.add_paragraph(
    "The calculations in this report are reproduced by the following scripts in the MidTerm directory:"
)
for filename, description in (
    ("question1.py", "Question 1 pairwise estimation, PSD test, Higham repair, and covariance construction."),
    ("question2.py", "Question 2 market-model estimation and analytical Normal VaR/ES."),
    ("question3.py", "Question 3 marginal fitting, copula comparison, simulation, and portfolio risk aggregation."),
):
    paragraph = document.add_paragraph(style="List Bullet")
    paragraph.add_run(filename).bold = True
    paragraph.add_run(f" — {description}")

document.add_paragraph(
    "Question 3 uses 50,000 simulations and random seed 545, so the reported Monte Carlo results are reproducible."
)

document.save(DOCX_FILE)
print(DOCX_FILE)
